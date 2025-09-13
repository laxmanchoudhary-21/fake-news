"""
Data sources module for fetching and verifying information from multiple APIs and feeds.
"""

import asyncio
import aiohttp
import feedparser
import requests
import json
import logging
from datetime import datetime, timedelta
from typing import Dict, List, Optional, Any, Tuple
from urllib.parse import urljoin, urlparse
import re

# Document processing
try:
    from newspaper import Article
    import PyPDF2
    from docx import Document
    from bs4 import BeautifulSoup
    DOCUMENT_PROCESSING = True
except ImportError:
    DOCUMENT_PROCESSING = False

from .config import DATA_SOURCES, settings
from .utils import cache_result, clean_text, extract_keywords, retry_on_failure, logger

class DataSourceManager:
    """Manage multiple data sources for verification and fact-checking."""
    
    def __init__(self):
        self.session = None
        self.rate_limiter = {}
        self._setup_session()
    
    def _setup_session(self):
        """Setup HTTP session with proper headers."""
        self.session = requests.Session()
        self.session.headers.update({
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36',
            'Accept': 'application/json, text/html, */*',
            'Accept-Language': 'en-US,en;q=0.9',
            'Accept-Encoding': 'gzip, deflate',
            'Connection': 'keep-alive',
        })
    
    async def verify_text_comprehensive(self, text: str, depth: str = "standard") -> Dict[str, Any]:
        """
        Comprehensive text verification across multiple sources.
        
        Args:
            text: Text to verify
            depth: Analysis depth (standard, deep, comprehensive)
            
        Returns:
            Comprehensive verification results
        """
        keywords = extract_keywords(text, 5)
        
        # Determine number of sources based on depth
        source_limits = {
            "standard": {"news": 5, "rss": 3, "fact_check": 2},
            "deep": {"news": 10, "rss": 5, "fact_check": 5},
            "comprehensive": {"news": 20, "rss": 10, "fact_check": 10}
        }
        limits = source_limits.get(depth, source_limits["standard"])
        
        # Create async tasks for all sources
        tasks = []
        
        # News APIs (if available)
        if settings.newsapi_key:
            tasks.append(self._fetch_newsapi(keywords, limits["news"]))
        if settings.newsdata_key:
            tasks.append(self._fetch_newsdata(keywords, limits["news"]))
        
        # Fact checking (if available)
        if settings.google_factcheck_key:
            tasks.append(self._fetch_google_factcheck(keywords, limits["fact_check"]))
        
        # RSS feeds (always available)
        tasks.append(self._fetch_rss_feeds(keywords, limits["rss"]))
        
        # Free alternative sources
        tasks.append(self._fetch_alternative_sources(keywords))
        
        # Execute all tasks concurrently
        try:
            results = await asyncio.gather(*tasks, return_exceptions=True)
            
            # Process results
            verification_data = self._process_verification_results(results, text)
            
            return verification_data
            
        except Exception as e:
            logger.error(f"Comprehensive verification failed: {e}")
            return {"error": str(e), "sources_checked": 0}
    
    @cache_result("newsapi", ttl=1800)  # 30 minutes cache
    async def _fetch_newsapi(self, keywords: List[str], limit: int) -> Dict[str, Any]:
        """Fetch results from NewsAPI."""
        if not settings.newsapi_key or not keywords:
            return {"source": "newsapi", "error": "API key not available or no keywords"}
        
        try:
            query = " ".join(keywords[:3])  # Use top 3 keywords
            
            async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=30)) as session:
                url = "https://newsapi.org/v2/everything"
                params = {
                    'q': query,
                    'apiKey': settings.newsapi_key,
                    'language': 'en',
                    'sortBy': 'relevancy',
                    'pageSize': min(limit, 50),
                    'from': (datetime.now() - timedelta(days=7)).strftime('%Y-%m-%d')
                }
                
                async with session.get(url, params=params) as response:
                    if response.status == 200:
                        data = await response.json()
                        if data.get('status') == 'ok':
                            articles = data.get('articles', [])
                            return {
                                "source": "newsapi",
                                "success": True,
                                "total_results": data.get('totalResults', 0),
                                "articles": [self._clean_article(article) for article in articles[:limit]],
                                "sources": list(set(article.get('source', {}).get('name', 'Unknown') for article in articles))
                            }
                    
                    return {"source": "newsapi", "error": f"API error: {response.status}"}
                    
        except Exception as e:
            return {"source": "newsapi", "error": str(e)}
    
    @cache_result("newsdata", ttl=1800)
    async def _fetch_newsdata(self, keywords: List[str], limit: int) -> Dict[str, Any]:
        """Fetch results from NewsData.io."""
        if not settings.newsdata_key or not keywords:
            return {"source": "newsdata", "error": "API key not available or no keywords"}
        
        try:
            query = keywords[0] if keywords else "news"  # Use primary keyword
            
            async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=30)) as session:
                url = "https://newsdata.io/api/1/news"
                params = {
                    'apikey': settings.newsdata_key,
                    'q': query,
                    'language': 'en',
                    'size': min(limit, 10)
                }
                
                async with session.get(url, params=params) as response:
                    if response.status == 200:
                        data = await response.json()
                        if data.get('status') == 'success':
                            articles = data.get('results', [])
                            return {
                                "source": "newsdata",
                                "success": True,
                                "total_results": len(articles),
                                "articles": [self._clean_article(article) for article in articles],
                                "sources": list(set(article.get('source_id', 'Unknown') for article in articles))
                            }
                    
                    return {"source": "newsdata", "error": f"API error: {response.status}"}
                    
        except Exception as e:
            return {"source": "newsdata", "error": str(e)}
    
    @cache_result("google_factcheck", ttl=3600)
    async def _fetch_google_factcheck(self, keywords: List[str], limit: int) -> Dict[str, Any]:
        """Fetch results from Google Fact Check Tools."""
        if not settings.google_factcheck_key or not keywords:
            return {"source": "google_factcheck", "error": "API key not available or no keywords"}
        
        try:
            query = " ".join(keywords[:2])  # Use top 2 keywords
            
            async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=30)) as session:
                url = "https://factchecktools.googleapis.com/v1alpha1/claims:search"
                params = {
                    'key': settings.google_factcheck_key,
                    'query': query,
                    'languageCode': 'en',
                    'maxAgeDays': 365
                }
                
                async with session.get(url, params=params) as response:
                    if response.status == 200:
                        data = await response.json()
                        claims = data.get('claims', [])
                        return {
                            "source": "google_factcheck",
                            "success": True,
                            "total_results": len(claims),
                            "claims": claims[:limit],
                            "fact_checkers": list(set([
                                review.get('publisher', {}).get('name', 'Unknown')
                                for claim in claims
                                for review in claim.get('claimReview', [])
                            ]))
                        }
                    
                    return {"source": "google_factcheck", "error": f"API error: {response.status}"}
                    
        except Exception as e:
            return {"source": "google_factcheck", "error": str(e)}
    
    @cache_result("rss_feeds", ttl=1800)
    async def _fetch_rss_feeds(self, keywords: List[str], max_per_feed: int) -> Dict[str, Any]:
        """Fetch results from RSS feeds."""
        if not keywords:
            return {"source": "rss", "error": "No keywords provided"}
        
        feeds = DATA_SOURCES["rss_feeds"]
        all_articles = []
        working_sources = []
        
        for source_name, feed_url in feeds:
            try:
                # Parse RSS feed
                feed = feedparser.parse(feed_url)
                
                if not feed.entries:
                    continue
                
                source_articles = []
                keywords_lower = [k.lower() for k in keywords]
                
                for entry in feed.entries[:20]:  # Check more entries per feed
                    title = getattr(entry, 'title', '').lower()
                    summary = getattr(entry, 'summary', '').lower()
                    description = getattr(entry, 'description', '').lower()
                    
                    # Check if any keyword appears in content
                    content_text = f"{title} {summary} {description}"
                    if any(keyword in content_text for keyword in keywords_lower):
                        article = {
                            "title": getattr(entry, 'title', 'No title'),
                            "link": getattr(entry, 'link', '#'),
                            "source": source_name,
                            "published": getattr(entry, 'published', 'Unknown date'),
                            "summary": (summary or description)[:300] + "..." if (summary or description) else "No summary"
                        }
                        source_articles.append(article)
                        
                        if len(source_articles) >= max_per_feed:
                            break
                
                if source_articles:
                    all_articles.extend(source_articles)
                    working_sources.append(source_name)
                    
            except Exception as e:
                logger.warning(f"RSS feed {source_name} failed: {e}")
                continue
        
        return {
            "source": "rss",
            "success": True if all_articles else False,
            "total_results": len(all_articles),
            "articles": all_articles[:50],  # Limit total RSS results
            "sources": working_sources
        }
    
    @cache_result("alternative_sources", ttl=3600)
    async def _fetch_alternative_sources(self, keywords: List[str]) -> Dict[str, Any]:
        """Fetch from free alternative sources when APIs are not available."""
        if not keywords:
            return {"source": "alternative", "error": "No keywords provided"}
        
        # Wikipedia search for factual information
        results = []
        
        try:
            # Simple Wikipedia API search
            wiki_query = keywords[0] if keywords else ""
            wiki_url = f"https://en.wikipedia.org/api/rest_v1/page/summary/{wiki_query}"
            
            async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=15)) as session:
                async with session.get(wiki_url) as response:
                    if response.status == 200:
                        data = await response.json()
                        if 'extract' in data:
                            results.append({
                                "title": data.get('title', 'Wikipedia Article'),
                                "summary": data.get('extract', '')[:500] + "...",
                                "source": "Wikipedia",
                                "link": data.get('content_urls', {}).get('desktop', {}).get('page', '#'),
                                "type": "encyclopedia"
                            })
            
            # Add more free sources here (academic papers, government sites, etc.)
            
        except Exception as e:
            logger.warning(f"Alternative sources search failed: {e}")
        
        return {
            "source": "alternative",
            "success": True if results else False,
            "total_results": len(results),
            "articles": results,
            "sources": ["Wikipedia"] if results else []
        }
    
    def _clean_article(self, article: Dict[str, Any]) -> Dict[str, Any]:
        """Clean and normalize article data."""
        return {
            "title": article.get('title', 'No title')[:200],
            "description": article.get('description', article.get('content', ''))[:500],
            "url": article.get('url', article.get('link', '#')),
            "source": article.get('source', {}).get('name') if isinstance(article.get('source'), dict) else article.get('source_id', 'Unknown'),
            "published_at": article.get('publishedAt', article.get('pubDate', 'Unknown date')),
            "author": article.get('author', 'Unknown author')
        }
    
    def _process_verification_results(self, results: List[Any], original_text: str) -> Dict[str, Any]:
        """Process and combine verification results from all sources."""
        verification_data = {
            "original_text": original_text,
            "sources_checked": 0,
            "total_articles": 0,
            "successful_sources": [],
            "failed_sources": [],
            "articles_by_source": {},
            "fact_checks": [],
            "verification_score": 0.0,
            "confidence_level": "low",
            "summary": ""
        }
        
        # Process each result
        for result in results:
            if isinstance(result, Exception):
                verification_data["failed_sources"].append(str(result))
                continue
            
            if not isinstance(result, dict):
                continue
            
            source_name = result.get("source", "unknown")
            verification_data["sources_checked"] += 1
            
            if result.get("success"):
                verification_data["successful_sources"].append(source_name)
                
                # Add articles
                articles = result.get("articles", [])
                if articles:
                    verification_data["articles_by_source"][source_name] = articles
                    verification_data["total_articles"] += len(articles)
                
                # Add fact checks
                claims = result.get("claims", [])
                if claims:
                    verification_data["fact_checks"].extend(claims)
            else:
                verification_data["failed_sources"].append(f"{source_name}: {result.get('error', 'Unknown error')}")
        
        # Calculate verification score
        verification_data["verification_score"] = self._calculate_verification_score(verification_data)
        
        # Determine confidence level
        score = verification_data["verification_score"]
        if score >= 0.7:
            verification_data["confidence_level"] = "high"
        elif score >= 0.4:
            verification_data["confidence_level"] = "medium"
        else:
            verification_data["confidence_level"] = "low"
        
        # Generate summary
        verification_data["summary"] = self._generate_verification_summary(verification_data)
        
        return verification_data
    
    def _calculate_verification_score(self, data: Dict[str, Any]) -> float:
        """Calculate overall verification score."""
        score = 0.0
        weights = {
            "sources_found": 0.3,
            "articles_found": 0.3,
            "fact_checks": 0.4
        }
        
        # Source diversity score
        num_sources = len(data["successful_sources"])
        source_score = min(num_sources / 5, 1.0)  # Normalize to 5 sources max
        
        # Article coverage score
        num_articles = data["total_articles"]
        article_score = min(num_articles / 20, 1.0)  # Normalize to 20 articles max
        
        # Fact check score
        num_fact_checks = len(data["fact_checks"])
        fact_check_score = min(num_fact_checks / 5, 1.0)  # Normalize to 5 fact checks max
        
        # Weighted combination
        score = (
            weights["sources_found"] * source_score +
            weights["articles_found"] * article_score +
            weights["fact_checks"] * fact_check_score
        )
        
        return score
    
    def _generate_verification_summary(self, data: Dict[str, Any]) -> str:
        """Generate human-readable verification summary."""
        successful = len(data["successful_sources"])
        total_checked = data["sources_checked"]
        articles = data["total_articles"]
        fact_checks = len(data["fact_checks"])
        
        summary = f"Verified across {successful}/{total_checked} sources. "
        
        if articles > 0:
            summary += f"Found {articles} related articles. "
        
        if fact_checks > 0:
            summary += f"Located {fact_checks} fact-check results. "
        
        if data["confidence_level"] == "high":
            summary += "High confidence in verification results."
        elif data["confidence_level"] == "medium":
            summary += "Moderate confidence in verification results."
        else:
            summary += "Limited verification data available."
        
        return summary
    
    @cache_result("url_extraction", ttl=3600)
    def extract_article_from_url(self, url: str) -> Dict[str, Any]:
        """
        Extract article content from URL.
        
        Args:
            url: Article URL
            
        Returns:
            Extracted article data
        """
        if not DOCUMENT_PROCESSING:
            return {"error": "Document processing not available. Install newspaper3k"}
        
        try:
            article = Article(url)
            article.download()
            article.parse()
            
            return {
                "success": True,
                "title": article.title,
                "text": article.text,
                "authors": article.authors,
                "publish_date": article.publish_date.isoformat() if article.publish_date else None,
                "top_image": article.top_image,
                "url": url,
                "summary": article.summary if hasattr(article, 'summary') else article.text[:500] + "..."
            }
            
        except Exception as e:
            return {"error": f"Failed to extract article: {str(e)}"}
    
    def extract_text_from_file(self, file_content: bytes, file_type: str) -> Dict[str, Any]:
        """
        Extract text from uploaded files.
        
        Args:
            file_content: File content as bytes
            file_type: File type (pdf, docx, txt)
            
        Returns:
            Extracted text data
        """
        if not DOCUMENT_PROCESSING:
            return {"error": "Document processing not available"}
        
        try:
            if file_type == "pdf":
                return self._extract_from_pdf(file_content)
            elif file_type == "docx":
                return self._extract_from_docx(file_content)
            elif file_type == "txt":
                return {"success": True, "text": file_content.decode('utf-8')}
            else:
                return {"error": f"Unsupported file type: {file_type}"}
                
        except Exception as e:
            return {"error": f"Text extraction failed: {str(e)}"}
    
    def _extract_from_pdf(self, file_content: bytes) -> Dict[str, Any]:
        """Extract text from PDF file."""
        try:
            import io
            pdf_file = io.BytesIO(file_content)
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            
            text = ""
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
            
            return {
                "success": True,
                "text": text,
                "pages": len(pdf_reader.pages),
                "metadata": pdf_reader.metadata if hasattr(pdf_reader, 'metadata') else {}
            }
            
        except Exception as e:
            return {"error": f"PDF extraction failed: {str(e)}"}
    
    def _extract_from_docx(self, file_content: bytes) -> Dict[str, Any]:
        """Extract text from DOCX file."""
        try:
            import io
            docx_file = io.BytesIO(file_content)
            doc = Document(docx_file)
            
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            return {
                "success": True,
                "text": text,
                "paragraphs": len(doc.paragraphs)
            }
            
        except Exception as e:
            return {"error": f"DOCX extraction failed: {str(e)}"}

# Global data source manager instance
_data_source_manager = None

def get_data_source_manager() -> DataSourceManager:
    """Get or create global data source manager instance."""
    global _data_source_manager
    if _data_source_manager is None:
        _data_source_manager = DataSourceManager()
    return _data_source_manager